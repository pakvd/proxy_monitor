from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook

HEADER_ALIASES = {
    "ip": "host",
    "host": "host",
    "хост": "host",
    "port": "port",
    "порт": "port",
    "login": "login",
    "логин": "login",
    "user": "login",
    "username": "login",
    "password": "password",
    "пароль": "password",
    "pass": "password",
    "address": "address",
    "адрес": "address",
    "server": "address",
    "сервер": "address",
    "name": "name",
    "имя": "name",
    "vlan": "name",
    "модем": "name",
    "protocol": "protocol",
    "протокол": "protocol",
}

HOST_RE = re.compile(r"^[A-Za-z0-9._:-]{1,253}$")
PROTOCOLS = {"http", "https", "socks5", "socks4"}


def parse_xlsx(raw: bytes) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if not raw:
        raise ValueError("файл пустой")
    try:
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("нужен файл .xlsx") from exc
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            raise ValueError("в файле нет строк")
        mapping = _map_header(header)
        if "host" not in mapping or "port" not in mapping:
            raise ValueError("в первой строке нужны колонки ip и port")
        parsed: list[dict[str, Any]] = []
        errors: list[dict[str, Any]] = []
        seen: set[tuple[str, int, str]] = set()
        for offset, values in enumerate(rows, start=2):
            if values is None or all(_blank(cell) for cell in values):
                continue
            row, error = _parse_row(values, mapping)
            if error:
                errors.append({"row": offset, "message": error})
                continue
            if row is None:
                continue
            key = (row["host"], row["port"], row["login"])
            if key in seen:
                errors.append({"row": offset, "message": "повтор ip, порта и логина в файле"})
                continue
            seen.add(key)
            parsed.append(row)
        return parsed, errors
    finally:
        workbook.close()


def workbook_bytes(rows: list[dict[str, Any]], *, example: bool = False) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "proxies"
    headers = ["ip", "port", "address", "name", "protocol"]
    if example:
        headers = ["ip", "port", "login", "password", "address", "name", "protocol"]
    sheet.append(headers)
    if example and not rows:
        sheet.append(["203.0.113.10", 30001, "", "", "farm-01", "vlan12", "http"])
    for row in rows:
        if example:
            sheet.append(
                [
                    row.get("host") or row.get("ip") or "",
                    row.get("port") or "",
                    "",
                    "",
                    row.get("address") or "",
                    row.get("name") or "",
                    row.get("protocol") or "http",
                ]
            )
            continue
        sheet.append(
            [
                row.get("host") or row.get("ip") or "",
                row.get("port") or "",
                row.get("address") or "",
                row.get("name") or "",
                row.get("protocol") or "http",
            ]
        )
    widths = (18, 10, 18, 18, 22, 16, 12)
    for column, width in zip("ABCDEFG", widths[: len(headers)]):
        sheet.column_dimensions[column].width = width
    last = chr(ord("A") + len(headers) - 1)
    sheet.auto_filter.ref = f"A1:{last}1"
    sheet.freeze_panes = "A2"
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _map_header(header: tuple[Any, ...] | None) -> dict[str, int]:
    mapping: dict[str, int] = {}
    if not header:
        return mapping
    for index, cell in enumerate(header):
        key = HEADER_ALIASES.get(_norm(cell))
        if key and key not in mapping:
            mapping[key] = index
    return mapping


def _parse_row(values: tuple[Any, ...], mapping: dict[str, int]) -> tuple[dict[str, Any] | None, str | None]:
    host = _cell(values, mapping.get("host"))
    port_raw = _cell(values, mapping.get("port"))
    if not host and not port_raw:
        return None, None
    if not host:
        return None, "не указан ip"
    if not port_raw:
        return None, "не указан порт"
    host = host.strip()
    if not HOST_RE.match(host):
        return None, "некорректный ip или хост"
    try:
        port = int(float(port_raw)) if port_raw else 0
    except (TypeError, ValueError):
        return None, "некорректный порт"
    if port < 1 or port > 65535:
        return None, "порт должен быть от 1 до 65535"
    login = _cell(values, mapping.get("login")).strip()
    password = _cell(values, mapping.get("password"))
    address = _cell(values, mapping.get("address")).strip()
    name = _cell(values, mapping.get("name")).strip()
    protocol = _cell(values, mapping.get("protocol")).strip().lower() or "http"
    if protocol not in PROTOCOLS:
        return None, "протокол должен быть http, https, socks5 или socks4"
    if len(login) > 200 or len(password) > 200 or len(address) > 200 or len(name) > 200:
        return None, "слишком длинное значение"
    return {
        "host": host,
        "port": port,
        "login": login,
        "password": password,
        "address": address,
        "name": name,
        "protocol": protocol,
    }, None


def _cell(values: tuple[Any, ...], index: int | None) -> str:
    if index is None or index >= len(values) or values[index] is None:
        return ""
    value = values[index]
    if isinstance(value, bool):
        return str(value).strip()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).strip()
    return str(value).strip()


def _blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def _norm(value: Any) -> str:
    return str(value or "").strip().lower().replace("ё", "е")
